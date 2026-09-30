#!/usr/bin/env python3
"""Junta as planilhas baixadas dos celulares em uma só, sem repetidos.

Uso: python3 juntar_planilhas.py entrada1.xlsx entrada2.xlsx [...] [-o saida.xlsx]

Quando o mesmo Nº USP aparece em mais de um arquivo (ou mais de uma vez),
fica a leitura mais antiga; o nome é aproveitado de qualquer leitura que o tenha.
A coluna "Leituras" mostra quantas vezes o número apareceu no total.
"""
import argparse
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

FMT = "%d/%m/%Y %H:%M:%S"


def ler(caminho):
    ws = openpyxl.load_workbook(caminho).active
    linhas = list(ws.iter_rows(values_only=True))
    cab = [str(c).strip() if c else "" for c in linhas[0]]
    for r in linhas[1:]:
        d = dict(zip(cab, r))
        if d.get("Nº USP"):
            yield {"Evento": d.get("Evento") or "", "numero": str(d["Nº USP"]), "nome": d.get("Nome") or "",
                   "hora": str(d.get("Data/hora") or ""), "aparelho": d.get("Aparelho") or Path(caminho).stem}


def quando(h):
    try:
        return datetime.strptime(h, FMT)
    except ValueError:
        return datetime.max


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("arquivos", nargs="+")
    ap.add_argument("-o", "--saida", default="presenca_unificada.xlsx")
    a = ap.parse_args()

    por_numero, total = {}, 0
    for arq in a.arquivos:
        for r in ler(arq):
            total += 1
            atual = por_numero.get(r["numero"])
            if atual is None:
                por_numero[r["numero"]] = {**r, "leituras": 1}
                continue
            atual["leituras"] += 1
            if not atual["nome"]:
                atual["nome"] = r["nome"]
            if quando(r["hora"]) < quando(atual["hora"]):
                atual.update(hora=r["hora"], aparelho=r["aparelho"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Presença"
    ws.append(["Evento", "Nº USP", "Nome", "Data/hora", "Aparelho", "Leituras"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for r in sorted(por_numero.values(), key=lambda r: quando(r["hora"])):
        ws.append([r["Evento"], r["numero"], r["nome"], r["hora"], r["aparelho"], r["leituras"]])
        ws.cell(ws.max_row, 2).number_format = "@"
    for col, larg in zip("ABCDEF", (24, 12, 40, 20, 16, 9)):
        ws.column_dimensions[col].width = larg
    ws.freeze_panes = "A2"
    wb.save(a.saida)
    print(f"{total} leituras -> {len(por_numero)} pessoas distintas ({total - len(por_numero)} repetidas). Salvo em {a.saida}")


if __name__ == "__main__":
    main()
